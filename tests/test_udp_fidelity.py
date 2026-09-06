"""
Tests for app/validation/udp_fidelity.py.

Self-contained: synthetic PowerDesigner and erwin XML files are written to a
temp directory so the MATCHED / MISSING / MISMATCH branches, the placeholder
filter, the multi-line value grammar and the fidelity blend are all exercised
without any of the large sample models.
"""

import os
import sys
from types import SimpleNamespace

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app.validation import udp_fidelity as uf  # noqa: E402


PD_CDM = """<?xml version="1.0" encoding="UTF-8"?>
<?PowerDesigner AppLocale="UTF16" ID="{1F5B9C98-0000-0000-0000-000000000001}" Label="" LastModificationDate="1" Name="Demo_CDM" Objects="4" Symbols="0" Target="" Type="{1E597170-9350-11D1-AB3C-0020AF71E433}" signature="CDM_DATA_MODEL_XML" version="16.7.6.7257"?>
<Model xmlns:a="attribute" xmlns:c="collection" xmlns:o="object">
<o:RootObject Id="o1"><c:Children>
<o:Model Id="o2"><a:Name>Demo_CDM</a:Name><a:Code>DEMO_CDM</a:Code>
<a:ExtendedAttributesText>{AAAA0000-0000-0000-0000-000000000001},Ext,60={AAAA0000-0000-0000-0000-000000000002},ModelOwner,4=Team
</a:ExtendedAttributesText>
<c:Entities>
<o:Entity Id="o3"><a:ObjectID>E1</a:ObjectID><a:Name>Customer</a:Name><a:Code>CUSTOMER</a:Code>
<a:ExtendedAttributesText>{AAAA0000-0000-0000-0000-000000000001},SHL_Ext,300={AAAA0000-0000-0000-0000-000000000003},Design_Status,5=Draft
{AAAA0000-0000-0000-0000-000000000004},Data_Category,13=&lt;unspecified&gt;
{AAAA0000-0000-0000-0000-000000000005},Maps_To,32="Sales Org"
"Legal Org"

{AAAA0000-0000-0000-0000-000000000006},Owner,5=Alice

</a:ExtendedAttributesText>
<c:Attributes>
<o:EntityAttribute Id="o4"><a:Name>Customer Id</a:Name><a:Code>CUSTOMER_ID</a:Code>
<a:ExtendedAttributesText>{AAAA0000-0000-0000-0000-000000000001},SHL_Ext,60={AAAA0000-0000-0000-0000-000000000007},PII,5=false
</a:ExtendedAttributesText>
</o:EntityAttribute>
</c:Attributes>
</o:Entity>
<o:Entity Id="o5"><a:Name>Dropped</a:Name><a:Code>DROPPED</a:Code>
<a:ExtendedAttributesText>{AAAA0000-0000-0000-0000-000000000001},SHL_Ext,60={AAAA0000-0000-0000-0000-000000000003},Design_Status,5=Final
</a:ExtendedAttributesText>
</o:Entity>
</c:Entities>
<c:Relationships><o:Relationship Id="o6"><a:Name>R1</a:Name>
<c:Object1><o:Entity Ref="o3"/></c:Object1></o:Relationship></c:Relationships>
</o:Model>
</c:Children></o:RootObject>
</Model>
"""

ERWIN_XML = """<?xml version="1.0" encoding="UTF-8"?>
<erwin xmlns="http://www.erwin.com/dm" xmlns:UDP="http://www.erwin.com/dm/metadata"
       xmlns:EMX="http://www.erwin.com/dm/data" FileVersion="10.10.38485" Format="erwin">
<EMX:Model id="{M}+00000001" name="Demo_CDM" xmlns="http://www.erwin.com/dm/data">
<ModelProps><Name>Demo_CDM</Name><UDP:ModelOwner>Team</UDP:ModelOwner></ModelProps>
<Udp_Groups>
  <Udp id="{U1}+00000000" name="Entity.Design_Status"><UdpProps><Name>Design_Status</Name><Long_Id>{U1}+00000000</Long_Id><Udp_Id>1</Udp_Id></UdpProps></Udp>
  <Udp id="{U2}+00000000" name="Entity.Maps_To"><UdpProps><Name>Maps_To</Name><Long_Id>{U2}+00000000</Long_Id></UdpProps></Udp>
</Udp_Groups>
<Entity_Groups>
<Entity id="{E1}+00000000" name="Customer">
  <EntityProps><Name>Customer</Name><Physical_Name>CUSTOMER</Physical_Name>
    <UDP:Entity.Design_Status>Draft</UDP:Entity.Design_Status>
    <UDP:_x007B_U2_x007D__x002B_00000000>"Sales Org"
"Legal Org"</UDP:_x007B_U2_x007D__x002B_00000000>
    <UDP:Owner>Bob</UDP:Owner>
    <UDP:Stray>leftover</UDP:Stray>
  </EntityProps>
  <Attribute_Groups>
    <Attribute id="{A1}+00000000" name="Customer Id">
      <AttributeProps><Name>Customer Id</Name><Physical_Name>CUSTOMER_ID</Physical_Name>
        <UDP:PII>false</UDP:PII></AttributeProps>
    </Attribute>
  </Attribute_Groups>
</Entity>
</Entity_Groups>
</EMX:Model>
</erwin>
"""


@pytest.fixture
def files(tmp_path):
    pd_file = tmp_path / "Demo.cdm"
    erwin_file = tmp_path / "Demo.xml"
    pd_file.write_text(PD_CDM, encoding="utf-8")
    erwin_file.write_text(ERWIN_XML, encoding="utf-8")
    return str(pd_file), str(erwin_file)


def test_parse_extended_attributes_text_handles_nesting_and_multiline():
    text = ("{A},Ext,300={B},Design_Status,5=Draft\n"
            "{C},Maps_To,32=\"Sales Org\"\n\"Legal Org\"\n\n"
            "{D},Owner,5=Alice\n\n")
    parsed = uf.parse_extended_attributes_text(text)
    assert parsed == [
        ("Ext.Design_Status", "Design_Status", "Draft"),
        ("Ext.Maps_To", "Maps_To", '"Sales Org"\n"Legal Org"'),
        ("Ext.Owner", "Owner", "Alice"),
    ]


def test_pd_extraction_skips_placeholders_and_pointer_stubs(files):
    pd_file, _ = files
    values, stats, notes = uf.extract_pd_udps(pd_file)
    names = sorted((v.object_type, v.object_name, v.udp_name) for v in values)
    assert names == [
        ("ATTRIBUTE", "Customer Id", "PII"),
        ("ENTITY", "Customer", "Design_Status"),
        ("ENTITY", "Customer", "Maps_To"),
        ("ENTITY", "Customer", "Owner"),
        ("ENTITY", "Dropped", "Design_Status"),
        ("MODEL", "Demo_CDM", "ModelOwner"),
    ]
    assert stats["blank"] == 1          # <unspecified>
    assert stats["populated"] == 6
    attribute = next(v for v in values if v.object_type == "ATTRIBUTE")
    assert attribute.owner_name == "Customer" and attribute.owner_code == "CUSTOMER"


def test_erwin_extraction_resolves_names_from_definitions(files):
    _, erwin_file = files
    values, stats, _ = uf.extract_erwin_udps(erwin_file)
    assert stats["definitions"] == 2
    got = {(v.object_type, v.object_name, v.udp_name): v.value for v in values}
    assert got[("ENTITY", "Customer", "Design_Status")] == "Draft"
    assert got[("ENTITY", "Customer", "Maps_To")] == '"Sales Org"\n"Legal Org"'
    assert got[("ATTRIBUTE", "Customer Id", "PII")] == "false"
    assert got[("MODEL", "Demo_CDM", "ModelOwner")] == "Team"


def test_compare_classifies_matched_missing_mismatch(files):
    pd_file, erwin_file = files
    result = uf.calculate_udp_fidelity(pd_file, erwin_file, "CDM")
    assert result.error == ""
    by_key = {(r.object_name, r.udp_name): r for r in result.rows}
    assert by_key[("Customer", "Design_Status")].status == uf.ST_MATCHED
    assert by_key[("Customer", "Maps_To")].status == uf.ST_MATCHED       # multi-line value
    assert by_key[("Customer Id", "PII")].status == uf.ST_MATCHED
    assert by_key[("Demo_CDM", "ModelOwner")].status == uf.ST_MATCHED
    assert by_key[("Customer", "Owner")].status == uf.ST_MISMATCH
    assert by_key[("Customer", "Owner")].erwin_value == "Bob"
    assert by_key[("Dropped", "Design_Status")].status == uf.ST_MISSING
    assert "no UDP values on this object" in by_key[("Dropped", "Design_Status")].note
    assert (result.udp_total, result.udp_matched, result.udp_mismatch, result.udp_missing) == (6, 4, 1, 1)
    assert result.udp_fidelity_score == pytest.approx(66.67, abs=0.01)
    assert result.udp_extra_in_erwin == 1                                  # UDP:Stray
    assert result.udp_blank_in_pd == 1


def test_blend_and_apply_reduce_overall_score_but_not_promotion_gate(files):
    pd_file, erwin_file = files
    config = SimpleNamespace(UDP_FIDELITY_ENABLED=True, UDP_FIDELITY_WEIGHT=0.25,
                             UDP_FIDELITY_AFFECTS_PROMOTION_GATE=False,
                             FIDELITY_REVIEW_THRESHOLD=95.0)
    result = SimpleNamespace(pd_file=pd_file, erwin_file=erwin_file, status="PASS",
                             fidelity_score=100.0, fidelity_score_raw=100.0, needs_review=False)
    udp = uf.apply(result, "CDM", config)
    assert udp is not None and udp.udp_total == 6
    assert result.structural_fidelity_score == 100.0
    assert result.fidelity_score == pytest.approx(100 * 0.75 + 66.67 * 0.25, abs=0.01)
    assert result.fidelity_score_raw == 100.0        # gate untouched by default
    assert result.needs_review is True
    assert result.udp_missing == 1 and result.udp_matched == 4

    gated = SimpleNamespace(pd_file=pd_file, erwin_file=erwin_file, status="PASS",
                            fidelity_score=100.0, fidelity_score_raw=100.0, needs_review=False)
    config.UDP_FIDELITY_AFFECTS_PROMOTION_GATE = True
    uf.apply(gated, "PDM", config)
    assert gated.fidelity_score_raw < 100.0


def test_apply_is_inert_when_disabled_or_error_or_no_udps(files, tmp_path):
    pd_file, erwin_file = files
    result = SimpleNamespace(pd_file=pd_file, erwin_file=erwin_file, status="PASS",
                             fidelity_score=90.0)
    assert uf.apply(result, "CDM", SimpleNamespace(UDP_FIDELITY_ENABLED=False)) is None
    assert result.fidelity_score == 90.0

    error = SimpleNamespace(pd_file=pd_file, erwin_file=erwin_file, status="ERROR",
                            fidelity_score=0.0)
    assert uf.apply(error, "CDM") is None

    # missing erwin file → recorded, score untouched, nothing raised
    broken = SimpleNamespace(pd_file=pd_file, erwin_file=str(tmp_path / "nope.xml"),
                             status="PASS", fidelity_score=90.0)
    udp = uf.apply(broken, "CDM", SimpleNamespace(UDP_FIDELITY_ENABLED=True))
    assert udp.error and broken.fidelity_score == 90.0

    # no populated UDPs → n/a, score untouched
    empty_pd = tmp_path / "Empty.cdm"
    empty_pd.write_text(PD_CDM.split("<a:ExtendedAttributesText>")[0]
                        + "</o:Model></c:Children></o:RootObject></Model>", encoding="utf-8")
    plain = SimpleNamespace(pd_file=str(empty_pd), erwin_file=erwin_file, status="PASS",
                            fidelity_score=90.0)
    udp = uf.apply(plain, "LDM", SimpleNamespace(UDP_FIDELITY_ENABLED=True))
    assert udp.udp_total == 0 and udp.udp_fidelity_score is None
    assert plain.fidelity_score == 90.0


def test_report_helpers_and_sheets(files):
    from openpyxl import Workbook
    pd_file, erwin_file = files
    result = SimpleNamespace(pd_file=pd_file, erwin_file=erwin_file, status="PASS",
                             fidelity_score=100.0, needs_review=False)
    uf.apply(result, "CDM", SimpleNamespace(UDP_FIDELITY_ENABLED=True, UDP_FIDELITY_WEIGHT=0.2))
    assert len(uf.summary_values(result)) == len(uf.SUMMARY_HEADERS) == len(uf.SUMMARY_WIDTHS)
    assert uf.summary_values(result)[:4] == [6, 4, 1, 1]

    wb = Workbook()
    ws = wb.active
    uf.summary_totals(ws, 5, 3, [result])
    assert ws.cell(5, 3).value == 6 and ws.cell(5, 7).value == pytest.approx(66.67, abs=0.01)
    stats = dict(uf.dashboard_statistics([result]))
    assert stats["UDP values compared (SAP PD)"] == 6

    uf.build_sheets(wb, [result], tier_label="TEST")
    assert "UDP_FIDELITY" in wb.sheetnames and "UDP_DETAIL" in wb.sheetnames
    detail = wb["UDP_DETAIL"]
    statuses = {detail.cell(r, 10).value for r in range(2, detail.max_row + 1)}
    assert statuses == {uf.ST_MATCHED, uf.ST_MISSING, uf.ST_MISMATCH}


# ─── promotion gate bands ─────────────────────────────────────────────────────

def test_promotion_gate_bands(tmp_path, monkeypatch):
    from app.validation import promotion_gate as pg
    from app.config import settings
    monkeypatch.setattr(settings, "PROMOTION_FIDELITY_THRESHOLD", 90.0, raising=False)
    monkeypatch.setattr(settings, "REVIEW_FIDELITY_FLOOR", 60.0, raising=False)
    monkeypatch.setattr(settings, "PROMOTION_FIDELITY_BASIS", "overall", raising=False)
    # These assertions describe the legacy three-band gate, which is still
    # available; the default is now two bands (see tests/test_promotion_bands.py).
    monkeypatch.setattr(settings, "PROMOTION_BANDS", "three", raising=False)

    def res(score, status="WARN"):
        return SimpleNamespace(pd_file="x.cdm", pd_model="x", status=status, fidelity_score=score,
                               critical_count=1, warning_count=2)

    good = res(92.0)
    assert pg.apply(good) is True and good.status == "PASS" and good.promotion_note.startswith("PASS")
    held = res(73.8, "FAIL")
    assert pg.apply(held) is False and held.status == "WARN" and "HELD in 2_preprocessed" in held.promotion_note
    bad = res(59.9)
    assert pg.apply(bad) is False and bad.status == "FAIL" and bad.promotion_note.startswith("REJECTED")
    assert pg.band(res(60.0)) == pg.BAND_REVIEW and pg.band(res(90.0)) == pg.BAND_PASS
    err = res(0.0, "ERROR")
    assert pg.apply(err) is False and err.status == "ERROR"

    staged = tmp_path / "2_preprocessed" / "xml" / "m.xml"
    staged.parent.mkdir(parents=True)
    staged.write_text("<x/>")
    moved = pg.reject(bad, str(staged))
    assert not staged.exists() and moved == [str(tmp_path / "2_preprocessed" / "xml" / "rejected" / "m.xml")]

    # Two-band (the default): the WARN band disappears and the model is FAILED.
    monkeypatch.setattr(settings, "PROMOTION_BANDS", "two", raising=False)
    monkeypatch.setattr(settings, "MANUAL_REVIEW_DIR", str(tmp_path / "manual_review"),
                        raising=False)
    failed = res(73.8)
    assert pg.apply(failed) is False
    assert failed.status == pg.BAND_FAIL and "manual_review" in failed.promotion_note
    assert pg.band(res(60.0)) == pg.BAND_FAIL
