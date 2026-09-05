"""
Mismatch Highlighter  (PDM → ERwin validation report post-processor)
====================================================================
Takes the ``validation_report.xlsx`` that ``report_generator.py`` already
produces and marks up every *non-match* so a human can review thousands of
models quickly:

  1. SUMMARY sheet
       • every mismatch counter cell that is > 0 (Tbl / Col / FK Missing &
         Extra, CRITICAL, WARNING) is painted red / amber and given an Excel
         note (hover-comment) that spells out, in plain English, what did not
         match and where to look.
       • a new "Manual Check" column is appended to each model row with a one
         line verdict, e.g.  "REVIEW → 7 cols missing, 8 FKs missing, 9 critical".
         Sort / filter on it to jump straight to the models that need eyes.

  2. FINDINGS sheet
       • a "Match?" column is inserted marking each row MISMATCH, and the
         Message cell is tinted by severity so the not-matched rows stand out.

  3. MANUAL_REVIEW sheet  (new)
       • one row per genuine mismatch (CRITICAL + WARNING by default), each with
         a ready-to-read "Review Message" sentence — table, column, what is
         wrong, PD value vs ERwin value, and the action to take.  AutoFilter is
         on, so a reviewer can filter to one model at a time.  This is the sheet
         that makes 5,000-model manual validation practical, because the
         per-model detail tabs are not emitted for large runs.

The script is header-driven (it locates columns by their header text, not a
fixed position), so it works on both the PDM and the CDM validator reports even
if the column order changes.

Usage
-----
  python highlight_mismatches.py                       # ./validation_report.xlsx  -> ./validation_report_highlighted.xlsx
  python highlight_mismatches.py path\to\report.xlsx   # explicit input
  python highlight_mismatches.py report.xlsx -o out.xlsx
  python highlight_mismatches.py report.xlsx --in-place # overwrite the input file
  python highlight_mismatches.py report.xlsx --include-info   # also flag INFO-level notes

Or call it from your batch runner right after generate_report():
  from highlight_mismatches import highlight_report
  highlight_report(report_path)          # writes <name>_highlighted.xlsx next to it

Dependencies:  pip install openpyxl
"""

import argparse
import os
import sys
import logging

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# ─── COLOUR PALETTE  (matches report_generator.py) ──────────────────────────────
C_RED     = "FFFF0000"   # missing / critical  — hard problem
C_AMBER   = "FFFFC000"   # extra / warning     — softer problem
C_BLUE    = "FF4472C4"   # info
C_GREEN   = "FF92D050"   # OK / matched
C_HEADER  = "FF1F3864"
C_WHITE   = "FFFFFFFF"
C_BLACK   = "FF000000"
C_OKTEXT  = "FF375623"

FILL_RED   = PatternFill("solid", fgColor=C_RED)
FILL_AMBER = PatternFill("solid", fgColor=C_AMBER)
FILL_BLUE  = PatternFill("solid", fgColor=C_BLUE)
FILL_GREEN = PatternFill("solid", fgColor=C_GREEN)


# Fix:
# The same nested severity->fill conditional appeared in two places
# (_style_match_marker and the manual-review row styling). Extracted here once,
# as plain if/elif/else. Same three fills, same order, same objects returned.
def _severity_fill(severity):
    """The cell fill for a finding severity: red, amber, or blue for the rest."""
    if severity == "CRITICAL":
        return FILL_RED
    if severity == "WARNING":
        return FILL_AMBER
    return FILL_BLUE

THIN   = Side(style="thin", color="FFB8B8B8")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FONT = Font(bold=True, color=C_WHITE, name="Calibri", size=10)
HEADER_FILL = PatternFill("solid", fgColor=C_HEADER)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)

COMMENT_AUTHOR = "Validator"

# ─── SUMMARY MISMATCH RULES ─────────────────────────────────────────────────────
# header text  ->  (fill, level, human phrase).  {n} is replaced with the value.
# level "hard" = red, "soft" = amber.  Only fired when the cell value is > 0.
SUMMARY_RULES = {
    "Tbl Missing": (FILL_RED,   "hard",
                    "{n} table(s) exist in PowerDesigner but are MISSING in ERwin."),
    "Tbl Extra":   (FILL_AMBER, "soft",
                    "{n} table(s) exist in ERwin but are NOT in PowerDesigner (extra)."),
    "Col Missing": (FILL_RED,   "hard",
                    "{n} column(s) exist in PowerDesigner but are MISSING in ERwin."),
    "Col Extra":   (FILL_AMBER, "soft",
                    "{n} column(s) exist in ERwin but are NOT in PowerDesigner (extra)."),
    "FK Missing":  (FILL_RED,   "hard",
                    "{n} foreign key(s) in PowerDesigner are MISSING in ERwin."),
    "FK Extra":    (FILL_AMBER, "soft",
                    "{n} foreign key(s) in ERwin are NOT in PowerDesigner (extra)."),
    "CRITICAL":    (FILL_RED,   "hard",
                    "{n} CRITICAL difference(s) — this model FAILED and must be reviewed."),
    "WARNING":     (FILL_AMBER, "soft",
                    "{n} WARNING difference(s) — please review."),
}

# short label used in the "Manual Check" verdict column
SUMMARY_SHORT = {
    "Tbl Missing": "{n} tbl missing",
    "Tbl Extra":   "{n} tbl extra",
    "Col Missing": "{n} cols missing",
    "Col Extra":   "{n} cols extra",
    "FK Missing":  "{n} FKs missing",
    "FK Extra":    "{n} FKs extra",
    "CRITICAL":    "{n} critical",
    "WARNING":     "{n} warning",
}


# ─── HELPERS ────────────────────────────────────────────────────────────────────

def _num(value) -> float:
    """Best-effort numeric read of a cell value (blank / text -> 0)."""
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _find_header_row(ws, wanted, max_scan=5):
    """
    Return (header_row_index, {header_text: column_index}) for the first row
    (within the first `max_scan` rows) that contains any of the `wanted` labels.
    Header text is matched case-insensitively and trimmed.
    """
    wanted_lc = {w.lower() for w in wanted}
    for row in range(1, max_scan + 1):
        mapping = {}
        for col in range(1, ws.max_column + 1):
            val = ws.cell(row=row, column=col).value
            if isinstance(val, str) and val.strip():
                mapping[val.strip()] = col
        if any(h.lower() in wanted_lc for h in mapping):
            return row, mapping
    return None, {}


def _col_by_name(mapping, *names):
    """First matching column index for any of the given header names (case-insensitive)."""
    lower = {k.lower(): v for k, v in mapping.items()}
    for name in names:
        if name.lower() in lower:
            return lower[name.lower()]
    return None


def _note(cell, text):
    """Attach an Excel note (hover comment) to a cell, sized to fit."""
    comment = Comment(text, COMMENT_AUTHOR)
    comment.width = 300
    comment.height = 120
    cell.comment = comment


# ─── 1. SUMMARY SHEET ───────────────────────────────────────────────────────────

def _format_summary_issue(cell, header, n, model_name):
    """Highlight one non-zero summary counter and attach its review note."""
    fill, _level, phrase = SUMMARY_RULES[header]
    cell.fill = fill
    cell.font = Font(bold=True, color=C_WHITE if fill is FILL_RED else C_BLACK)
    cell.alignment = CENTER
    msg = phrase.format(n=int(n))
    if model_name:
        msg += (
            f"\nOpen the FINDINGS / MANUAL_REVIEW sheet and filter "
            f"Model = '{model_name}' to see exactly which."
        )
    _note(cell, msg)


def _process_summary_rule_cells(ws, row, rule_cols, model_name):
    """Process all mismatch counters for one SUMMARY row."""
    parts = []
    for header, col in rule_cols.items():
        cell = ws.cell(row=row, column=col)
        n = _num(cell.value)
        if n <= 0:
            continue
        _format_summary_issue(cell, header, n, model_name)
        parts.append(SUMMARY_SHORT[header].format(n=int(n)))
    return parts


def _set_summary_verdict(ws, row, verdict_col, parts):
    """Write the human-readable Manual Check verdict."""
    vcell = ws.cell(row=row, column=verdict_col)
    vcell.border = BORDER
    vcell.alignment = LEFT
    if parts:
        vcell.value = "REVIEW → " + ", ".join(parts)
        vcell.fill = FILL_AMBER
        vcell.font = Font(bold=True, color=C_BLACK)
        return True
    vcell.value = "OK — full match"
    vcell.font = Font(bold=True, color=C_OKTEXT)
    return False


def _flag_summary_status(ws, row, status_col):
    """Add a note to a non-PASS status cell."""
    if not status_col:
        return
    cell = ws.cell(row=row, column=status_col)
    value = cell.value
    if isinstance(value, str) and value.strip().upper() not in ("PASS", ""):
        _note(
            cell,
            f"Status = {value}. See the Manual Check column for a summary "
            f"and the MANUAL_REVIEW sheet for line-by-line detail.",
        )


def _is_summary_tail(ws, row, model_col):
    """Return True when the SUMMARY row is the total/blank tail."""
    first = ws.cell(row=row, column=1).value
    model_name = ws.cell(row=row, column=model_col).value if model_col else None
    if isinstance(first, str) and first.strip().upper() == "TOTAL":
        return True
    return first is None and model_name is None


def _extend_summary_filter(ws, verdict_col, end_row):
    """Extend the existing SUMMARY autofilter through Manual Check."""
    if not ws.auto_filter.ref:
        return
    try:
        last = get_column_letter(verdict_col)
        top_left = ws.auto_filter.ref.split(":")[0]
        ws.auto_filter.ref = f"{top_left}:{last}{end_row}"
    except Exception:  # nosec B110
        pass


def _process_summary(ws):
    """Highlight mismatch cells, add notes, and append a Manual Check verdict."""
    header_row, headers = _find_header_row(
        ws, list(SUMMARY_RULES) + ["Status", "PD File"]
    )
    if not header_row:
        logger.warning("SUMMARY: could not locate a header row; skipping.")
        return 0

    model_col = _col_by_name(headers, "PD File", "CDM File", "Model")
    status_col = _col_by_name(headers, "Status")
    rule_cols = {h: headers[h] for h in SUMMARY_RULES if h in headers}

    verdict_col = ws.max_column + 1
    hdr = ws.cell(row=header_row, column=verdict_col, value="Manual Check")
    hdr.font, hdr.fill, hdr.alignment, hdr.border = (
        HEADER_FONT, HEADER_FILL, CENTER, BORDER
    )
    ws.column_dimensions[get_column_letter(verdict_col)].width = 42

    flagged_rows = 0
    row = header_row + 1
    while row <= ws.max_row:
        if _is_summary_tail(ws, row, model_col):
            break
        model_name = (
            ws.cell(row=row, column=model_col).value if model_col else None
        )
        parts = _process_summary_rule_cells(ws, row, rule_cols, model_name)
        if _set_summary_verdict(ws, row, verdict_col, parts):
            flagged_rows += 1
        _flag_summary_status(ws, row, status_col)
        row += 1

    _extend_summary_filter(ws, verdict_col, row - 1)
    return flagged_rows


# ─── 2. FINDINGS SHEET  → build MANUAL_REVIEW from it ───────────────────────────

MISMATCH_LABEL = {
    "TABLE":       "Table not matched",
    "COLUMN":      "Column not matched",
    "DATA_TYPE":   "Data type differs",
    "NULLABILITY": "Nullability differs",
    "DEFAULT":     "Default value differs",
    "PRIMARY_KEY": "Primary key differs",
    "FOREIGN_KEY": "Foreign key not matched",
    "INDEX":       "Index not matched",
}


def _finding_columns(headers):
    """Resolve FINDINGS columns from their header names."""
    return {
        "model": _col_by_name(headers, "Model"),
        "status": _col_by_name(headers, "Status"),
        "fidelity": _col_by_name(headers, "Fidelity %"),
        "category": _col_by_name(headers, "Category"),
        "severity": _col_by_name(headers, "Severity"),
        "table": _col_by_name(headers, "Table", "Object"),
        "column": _col_by_name(headers, "Column", "Member"),
        "message": _col_by_name(headers, "Message"),
        "pd": _col_by_name(headers, "PD Value", "PowerDesigner Value"),
        "ew": _col_by_name(headers, "ERwin Value", "erwin Value"),
    }


def _shift_column(column, insert_at):
    """Shift a column index after inserting a new column."""
    return column + 1 if column is not None and column >= insert_at else column


def _shift_finding_columns(columns, insert_at):
    """Apply the FINDINGS column insertion offset."""
    return {
        name: _shift_column(column, insert_at)
        for name, column in columns.items()
    }


def _read_finding_cell(ws, row, column):
    """Read a FINDINGS cell, returning None for an unavailable column."""
    return ws.cell(row=row, column=column).value if column else None


def _finding_severity(ws, row, columns):
    """Read and normalize a finding severity."""
    value = _read_finding_cell(ws, row, columns["severity"])
    return value.strip().upper() if isinstance(value, str) else value


def _style_match_marker(ws, row, insert_at, severity):
    """Create and style the inserted Match? marker."""
    mark = ws.cell(row=row, column=insert_at, value="MISMATCH")
    mark.font = Font(
        bold=True,
        color=C_WHITE if severity == "CRITICAL" else C_BLACK,
    )
    mark.fill = _severity_fill(severity)
    mark.alignment = CENTER
    mark.border = BORDER


def _style_finding_message(ws, row, columns, severity):
    """Tint critical/warning message text."""
    if not columns["message"] or severity not in ("CRITICAL", "WARNING"):
        return
    cell = ws.cell(row=row, column=columns["message"])
    cell.font = Font(
        bold=severity == "CRITICAL",
        color=C_RED if severity == "CRITICAL" else "FFBF8F00",
    )


def _finding_row(ws, row, columns):
    """Convert one FINDINGS row into the normalized review structure."""
    return {
        "model": _read_finding_cell(ws, row, columns["model"]),
        "status": _read_finding_cell(ws, row, columns["status"]),
        "fidelity": _read_finding_cell(ws, row, columns["fidelity"]),
        "category": _read_finding_cell(ws, row, columns["category"]),
        "severity": _finding_severity(ws, row, columns),
        "table": _read_finding_cell(ws, row, columns["table"]),
        "column": _read_finding_cell(ws, row, columns["column"]),
        "message": _read_finding_cell(ws, row, columns["message"]),
        "pd": _read_finding_cell(ws, row, columns["pd"]),
        "ew": _read_finding_cell(ws, row, columns["ew"]),
    }


def _extend_findings_filter(ws, header_row):
    """Extend FINDINGS autofilter across the inserted Match? column."""
    if not ws.auto_filter.ref:
        return
    try:
        last = get_column_letter(ws.max_column)
        end_row = ws.auto_filter.ref.split(":")[-1]
        end_row = "".join(ch for ch in end_row if ch.isdigit())
        ws.auto_filter.ref = f"A{header_row}:{last}{end_row}"
    except Exception:  # nosec B110
        pass


def _process_findings(ws):
    """Mark every finding row as a MISMATCH and tint the message; return the rows."""
    header_row, headers = _find_header_row(ws, ["Message", "Severity", "Category"])
    if not header_row:
        logger.warning("FINDINGS: could not locate a header row; skipping.")
        return []

    columns = _finding_columns(headers)
    severity_col = columns["severity"]
    insert_at = severity_col + 1 if severity_col else 1
    ws.insert_cols(insert_at)
    columns = _shift_finding_columns(columns, insert_at)

    hdr = ws.cell(row=header_row, column=insert_at, value="Match?")
    hdr.font, hdr.fill, hdr.alignment, hdr.border = (
        HEADER_FONT, HEADER_FILL, CENTER, BORDER
    )
    ws.column_dimensions[get_column_letter(insert_at)].width = 12

    rows = []
    for row in range(header_row + 1, ws.max_row + 1):
        if (
            _read_finding_cell(ws, row, columns["message"]) is None
            and _read_finding_cell(ws, row, columns["category"]) is None
        ):
            continue
        severity = _finding_severity(ws, row, columns)
        _style_match_marker(ws, row, insert_at, severity)
        _style_finding_message(ws, row, columns, severity)
        rows.append(_finding_row(ws, row, columns))

    _extend_findings_filter(ws, header_row)
    return rows


# ─── 3. MANUAL_REVIEW SHEET ─────────────────────────────────────────────────────

MANUAL_HEADERS = [
    "#", "Model", "Status", "Fidelity %", "Severity", "Mismatch Type",
    "Table", "Column", "Review Message (what to check manually)",
    "PD Value", "ERwin Value",
]


def _manual_review_rows(findings, include_info):
    """Filter and sort findings that require manual review."""
    keep = ("CRITICAL", "WARNING") + (("INFO",) if include_info else ())
    order = {"CRITICAL": 0, "WARNING": 1, "INFO": 2}
    rows = [f for f in findings if (f["severity"] or "INFO") in keep]
    rows.sort(
        key=lambda f: (
            str(f["model"] or ""),
            order.get(f["severity"], 3),
            str(f["category"] or ""),
        )
    )
    return rows


def _manual_review_text(finding):
    """Build the ready-to-read manual review sentence."""
    table = finding["table"] or ""
    column = finding["column"] or ""
    where = ".".join(x for x in (table, column) if x) or "(model level)"
    msg = finding["message"] or MISMATCH_LABEL.get(
        finding["category"], "Difference"
    )
    pd_value = "" if finding["pd"] in (None, "—") else finding["pd"]
    erwin_value = "" if finding["ew"] in (None, "—") else finding["ew"]

    detail = ""
    if pd_value or erwin_value:
        detail = (
            f"  [PD: {pd_value or '(none)'}  vs  "
            f"ERwin: {erwin_value or '(none)'}]"
        )
    return f"{where} — {msg}.{detail}  → Verify manually in the model."


def _write_manual_review_row(ws, row, index, finding):
    """Write and style one MANUAL_REVIEW row."""
    values = [
        index,
        finding["model"],
        finding["status"],
        finding["fidelity"],
        finding["severity"],
        MISMATCH_LABEL.get(finding["category"], finding["category"]),
        finding["table"],
        finding["column"],
        _manual_review_text(finding),
        finding["pd"],
        finding["ew"],
    ]
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col, value=value)
        cell.alignment, cell.border = LEFT, BORDER

    sev_cell = ws.cell(row=row, column=5)
    severity = finding["severity"]
    sev_cell.fill = _severity_fill(severity)
    sev_cell.font = Font(
        bold=True,
        color=C_WHITE if severity in ("CRITICAL", "INFO") else C_BLACK,
    )
    sev_cell.alignment = CENTER


def _style_manual_review_sheet(ws):
    """Apply widths and filtering to MANUAL_REVIEW."""
    widths = [6, 30, 8, 10, 10, 22, 26, 22, 80, 26, 26]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.auto_filter.ref = (
        f"A1:{get_column_letter(len(MANUAL_HEADERS))}{max(ws.max_row, 1)}"
    )


def _build_manual_review(wb, findings, include_info=False):
    """Create a single flat sheet of only the mismatches, with a review sentence."""
    if "MANUAL_REVIEW" in wb.sheetnames:
        del wb["MANUAL_REVIEW"]
    ws = wb.create_sheet("MANUAL_REVIEW", index=1)
    ws.freeze_panes = "A2"

    for col, value in enumerate(MANUAL_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=value)
        cell.font, cell.fill, cell.alignment, cell.border = (
            HEADER_FONT, HEADER_FILL, CENTER, BORDER
        )

    rows = _manual_review_rows(findings, include_info)
    for index, finding in enumerate(rows, start=1):
        _write_manual_review_row(ws, index + 1, index, finding)

    if not rows:
        cell = ws.cell(
            row=2,
            column=1,
            value="No CRITICAL or WARNING mismatches found — all models match.",
        )
        cell.font = Font(bold=True, color=C_OKTEXT)

    _style_manual_review_sheet(ws)
    return len(rows)


# ─── PUBLIC API ─────────────────────────────────────────────────────────────────

def highlight_report(input_path, output_path=None, in_place=False, include_info=False):
    """
    Post-process a validation_report.xlsx: highlight non-matches, add notes,
    a Manual Check column, and a MANUAL_REVIEW sheet.  Returns the output path.
    """
    if not os.path.exists(input_path):
        raise FileNotFoundError(input_path)

    wb = load_workbook(input_path)

    flagged = 0
    if "SUMMARY" in wb.sheetnames:
        flagged = _process_summary(wb["SUMMARY"])
    else:
        logger.warning("No SUMMARY sheet found in %s", input_path)

    findings = []
    if "FINDINGS" in wb.sheetnames:
        findings = _process_findings(wb["FINDINGS"])
    else:
        logger.warning("No FINDINGS sheet found in %s", input_path)

    reviewed = _build_manual_review(wb, findings, include_info=include_info)

    if in_place:
        output_path = input_path
    elif not output_path:
        base, ext = os.path.splitext(input_path)
        output_path = f"{base}_highlighted{ext}"

    wb.save(output_path)
    logger.info("Highlighted report saved → %s", output_path)
    print(f"[OK] {flagged} model(s) flagged for review, "
          f"{reviewed} mismatch line(s) in MANUAL_REVIEW.")
    print(f"[OK] Saved → {output_path}")
    return output_path


def _parse_args():
    p = argparse.ArgumentParser(
        description="Highlight non-matches in a PDM→ERwin validation report.")
    p.add_argument("input", nargs="?", default="validation_report.xlsx",
                   help="Path to the validation report .xlsx (default: ./validation_report.xlsx)")
    p.add_argument("-o", "--output", help="Output path (default: <input>_highlighted.xlsx)")
    p.add_argument("--in-place", action="store_true", help="Overwrite the input file")
    p.add_argument("--include-info", action="store_true",
                   help="Also include INFO-level notes in MANUAL_REVIEW")
    return p.parse_args()


def main():
    logging.basicConfig(level=logging.INFO, format="%(levelname)-8s %(message)s")
    args = _parse_args()
    try:
        highlight_report(args.input, output_path=args.output,
                         in_place=args.in_place, include_info=args.include_info)
    except FileNotFoundError:
        print(f"[ERROR] File not found: {args.input}")
        print("        Run this from your report folder, or pass the path, e.g.:")
        print(r"        python highlight_mismatches.py C:\...\ValidationReports\validation_report.xlsx")
        sys.exit(1)


if __name__ == "__main__":
    main()