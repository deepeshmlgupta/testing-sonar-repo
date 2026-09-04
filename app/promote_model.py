import os
import shutil
from pathlib import Path
import openpyxl
import re

def run_promotion_scan():
    print("==================================================")
    print("HUMAN-IN-THE-LOOP PROMOTION ENGINE")
    print("==================================================\n")
    
    base_dir = Path(__file__).resolve().parent.parent
    manual_review_dir = base_dir / "manual_review_reports"
    preprocessed_dir = base_dir / "erwinmodels" / "2_preprocessed"
    final_dir = base_dir / "erwinmodels" / "3_final"
    
    os.makedirs(manual_review_dir, exist_ok=True)
    os.makedirs(final_dir / "erwin", exist_ok=True)
    os.makedirs(final_dir / "xml", exist_ok=True)
    
    excel_files = list(manual_review_dir.rglob("*.xlsx"))
    
    if not excel_files:
        print(f"No Excel reports found in '{manual_review_dir.name}'.")
        print("Drop a reviewed Excel report in this folder to promote it.")
        return
        
    for excel_path in excel_files:
        _process_excel_report(excel_path, manual_review_dir, preprocessed_dir, final_dir)

def _process_excel_report(excel_path, manual_review_dir, preprocessed_dir, final_dir):
    raw_stem = excel_path.stem
    base_model_name = re.sub(r'_\d+\.\d+%_Fidelity$', '', raw_stem)
    
    print(f"\nProcessing report: {excel_path.name}")
    print(f"Target model: {base_model_name}")
        
    try:
        wb = openpyxl.load_workbook(excel_path)
    except Exception as e:
        print(f"  [ERROR] Could not read {excel_path.name}: {e}")
        return
        
    all_warnings_acceptable, found_mismatches = _evaluate_findings(wb)
            
    if not found_mismatches:
        print("  [INFO] No mismatch sheets found to review.")
        wb.close()
        return
        
    if all_warnings_acceptable:
        print(f"  -> All warnings marked 'Acceptable'! Promoting {base_model_name} to 3_final...")
        _promote_files(base_model_name, wb, excel_path, preprocessed_dir, final_dir, manual_review_dir)
    else:
        wb.close()

def _evaluate_findings(wb):
    all_warnings_acceptable = True
    found_mismatches = False
    
    for sheet_name in wb.sheetnames:
        if sheet_name == "FINDINGS":
            ws = wb[sheet_name]
            review_col_idx, severity_col_idx = _find_header_columns(ws)
                    
            if review_col_idx is None or severity_col_idx is None:
                print(f"  [WARNING] Required columns not found in sheet '{sheet_name}'. Skipping.")
                return False, False
                
            found_mismatches = True
            if not _scan_finding_rows(ws, sheet_name, review_col_idx, severity_col_idx):
                return False, True
                
    return all_warnings_acceptable, found_mismatches

def _find_header_columns(ws):
    header_row = ws[1]
    review_col_idx = None
    severity_col_idx = None
    for cell in header_row:
        if cell.value and "Manual Review" in str(cell.value):
            review_col_idx = cell.column
        elif cell.value and "Severity" in str(cell.value):
            severity_col_idx = cell.column
    return review_col_idx, severity_col_idx

def _scan_finding_rows(ws, sheet_name, review_col_idx, severity_col_idx):
    for row_idx in range(2, ws.max_row + 1):
        mapping_cell = ws.cell(row=row_idx, column=1).value
        if not mapping_cell:
            continue
            
        severity_val = str(ws.cell(row=row_idx, column=severity_col_idx).value).strip().upper()
        if severity_val == "INFO":
            continue
            
        if severity_val == "CRITICAL":
            print(f"  -> Found CRITICAL issue on row {row_idx} of '{sheet_name}'. Criticals cannot be overridden. Promotion blocked.")
            return False
            
        review_value = ws.cell(row=row_idx, column=review_col_idx).value
        if not review_value or str(review_value).strip().lower() != "acceptable":
            print(f"  -> Found unapproved WARNING on row {row_idx} of '{sheet_name}'. Promotion blocked.")
            return False
    return True

def _promote_files(base_model_name, wb, excel_path, preprocessed_dir, final_dir, manual_review_dir):
    src_erwin = preprocessed_dir / "erwin" / f"{base_model_name}.erwin"
    src_xml = preprocessed_dir / "xml" / f"{base_model_name}.xml"
    dest_erwin = final_dir / "erwin" / f"{base_model_name}.erwin"
    dest_xml = final_dir / "xml" / f"{base_model_name}.xml"
    
    promoted = False
    if src_erwin.exists():
        shutil.copy2(src_erwin, dest_erwin)
        print(f"    Copied {base_model_name}.erwin")
        promoted = True
    
    if src_xml.exists():
        shutil.copy2(src_xml, dest_xml)
        print(f"    Copied {base_model_name}.xml")
        promoted = True
        
    if promoted:
        print("  [SUCCESS] Promotion complete.")
        _update_and_archive_excel(wb, excel_path, base_model_name, manual_review_dir)
    else:
        print(f"  [ERROR] Source files missing in {preprocessed_dir}. Could not promote.")
        wb.close()

def _update_and_archive_excel(wb, excel_path, base_model_name, manual_review_dir):
    if "SUMMARY" in wb.sheetnames:
        ws_summary = wb["SUMMARY"]
        for cell in ws_summary[2]:
            if cell.value == "Fidelity %":
                for r in range(3, ws_summary.max_row + 1):
                    if ws_summary.cell(row=r, column=cell.column).value is not None:
                        ws_summary.cell(row=r, column=cell.column).value = "100.0 (Approved)"
                break
                
    archive_dir = manual_review_dir / "archived_promotions"
    archive_dir.mkdir(exist_ok=True)
    
    new_filename = f"{base_model_name}_100.0%_Fidelity_Approved.xlsx"
    new_filepath = archive_dir / new_filename
    
    wb.save(new_filepath)
    wb.close()
    excel_path.unlink()
    
    print(f"  -> Updated Excel score to 100% and moved to archive: {new_filename}")

if __name__ == "__main__":
    run_promotion_scan()
